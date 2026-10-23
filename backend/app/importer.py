"""Turn the files an SME already has into customer rows.

Supported:
  * vCard (.vcf)  - iPhone "Share contacts", Android Contacts export, Google Contacts export
  * CSV           - Google Contacts / Outlook exports, or any sheet with name + phone columns
  * Excel (.xlsx) - first sheet, same column rules as CSV
  * Pasted text   - one contact per line, e.g. "Ahmed 050 123 4567"

Nothing here touches the database; it only returns clean rows for the review screen.
"""

from __future__ import annotations

import csv
import io
import quopri
import re
from dataclasses import dataclass

from .utils import normalize_phone

MAX_ROWS = 5000
_PHONE_RE = re.compile(r"(\+?\d[\d\s\-().]{6,}\d)")
_EMAIL_RE = re.compile(r"[^@\s,;<>]+@[^@\s,;<>]+\.[A-Za-z]{2,}")


@dataclass
class Row:
    name: str
    phone: str | None = None
    email: str | None = None
    area: str | None = None


def clean_phone(raw: str | None) -> str | None:
    """'050 123 4567' -> '+971501234567'. Returns None for anything that isn't a plausible number."""
    n = normalize_phone(raw)
    if not n or not (8 <= len(n) <= 15):
        return None
    return "+" + n


def _clean(v: str | None, limit: int) -> str | None:
    v = (v or "").strip().strip('"').strip()
    return v[:limit] or None


def finish_rows(rows: list[Row]) -> list[Row]:
    out = []
    for r in rows:
        r.name = (_clean(r.name, 160) or "")
        r.email = _clean(r.email, 200) if r.email and _EMAIL_RE.fullmatch(r.email.strip()) else None
        r.area = _clean(r.area, 120)
        if not r.name and r.phone:
            r.name = r.phone
        if r.name or r.phone:
            out.append(r)
        if len(out) >= MAX_ROWS:
            break
    return out


# ---------------------------------------------------------------- vCard
def _unfold(text: str) -> list[str]:
    """vCard folds long lines; quoted-printable uses '=' at line end as a soft break."""
    lines: list[str] = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if raw[:1] in (" ", "\t") and lines:
            lines[-1] += raw[1:]
        elif lines and lines[-1].endswith("=") and "QUOTED-PRINTABLE" in lines[-1].upper():
            lines[-1] = lines[-1][:-1] + raw
        else:
            lines.append(raw)
    return lines


def _vvalue(params: str, value: str) -> str:
    p = params.upper()
    if "QUOTED-PRINTABLE" in p:
        charset = re.search(r"CHARSET=([\w-]+)", params, re.I)
        value = quopri.decodestring(value.encode("latin-1", "ignore")).decode(
            charset.group(1) if charset else "utf-8", "replace"
        )
    return value.replace("\\,", ",").replace("\\;", ";").replace("\\n", " ").replace("\\N", " ").strip()


def parse_vcf(text: str) -> list[Row]:
    rows: list[Row] = []
    cur: dict | None = None
    for line in _unfold(text):
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = re.sub(r"^item\d+\.", "", key, flags=re.I)  # iPhone groups: item1.TEL
        name, _, params = key.partition(";")
        name = name.upper()
        if name == "BEGIN" and value.strip().upper() == "VCARD":
            cur = {"fn": "", "n": "", "tels": [], "email": None, "adr": None}
        elif name == "END" and cur is not None:
            tels = sorted(cur["tels"], key=lambda t: 0 if re.search(r"CELL|MOBILE|WHATSAPP", t[0], re.I) else 1)
            phone = next((p for _, raw in tels if (p := clean_phone(raw))), None)
            full = cur["fn"] or " ".join(x for x in reversed(cur["n"].split(";")[:2]) if x).strip()
            rows.append(Row(name=full, phone=phone, email=cur["email"], area=cur["adr"]))
            cur = None
        elif cur is None:
            continue
        elif name == "FN":
            cur["fn"] = _vvalue(params, value)
        elif name == "N":
            cur["n"] = _vvalue(params, value)
        elif name == "TEL":
            cur["tels"].append((params, _vvalue(params, value)))
        elif name == "EMAIL" and not cur["email"]:
            cur["email"] = _vvalue(params, value)
        elif name == "ADR" and not cur["adr"]:
            # ADR = PO box; extended; street; city; region; postcode; country
            parts = (_vvalue(params, value).split(";") + [""] * 7)[:7]
            cur["adr"] = next((parts[i].strip() for i in (3, 4, 2) if parts[i].strip()), None)
    return finish_rows(rows)


# ---------------------------------------------------------------- tables (CSV / Excel)
def _pick(headers: list[str], *tests) -> int | None:
    for test in tests:
        for i, h in enumerate(headers):
            if test(h):
                return i
    return None


def parse_table(table: list[list[str]]) -> list[Row]:
    table = [[(c or "").strip() for c in r] for r in table if any((c or "").strip() for c in r)]
    if not table:
        return []
    h = [c.lower() for c in table[0]]
    is_phone = lambda x: any(k in x for k in ("mobile", "phone", "tel", "whatsapp", "number", "contact no"))  # noqa: E731
    phone_i = _pick(
        h,
        lambda x: ("mobile" in x or "whatsapp" in x) and "type" not in x,
        lambda x: is_phone(x) and "value" in x,  # Google: "Phone 1 - Value"
        lambda x: is_phone(x) and "type" not in x and "label" not in x,
    )
    name_i = _pick(h, lambda x: x in ("name", "full name", "display name", "customer", "customer name", "client", "client name", "contact name"))
    first_i = _pick(h, lambda x: x in ("first name", "given name", "first"))
    last_i = _pick(h, lambda x: x in ("last name", "family name", "surname", "last"))
    email_i = _pick(h, lambda x: ("e-mail" in x or "email" in x) and ("value" in x or "address" in x or x in ("email", "e-mail")))
    area_i = _pick(h, lambda x: x in ("area", "city", "location", "community", "emirate", "address", "home city", "business city"))

    has_header = phone_i is not None or name_i is not None or first_i is not None
    body = table[1:] if has_header else table
    if not has_header:  # no headers: find the phone column by content, name is the first other column
        sample = body[: min(20, len(body))]
        width = max(len(r) for r in sample)
        scores = [sum(1 for r in sample if i < len(r) and clean_phone(r[i])) for i in range(width)]
        phone_i = max(range(width), key=lambda i: scores[i]) if any(scores) else None
        name_i = next((i for i in range(width) if i != phone_i), None)

    def cell(r, i):
        return r[i] if i is not None and i < len(r) else ""

    rows = []
    for r in body:
        name = cell(r, name_i) or " ".join(x for x in (cell(r, first_i), cell(r, last_i)) if x)
        raw_phone = cell(r, phone_i)
        # Google puts several numbers in one cell: "050 1 ::: 04 2"
        phone = next((p for part in re.split(r":::|/|,", raw_phone) if (p := clean_phone(part))), None)
        rows.append(Row(name=name, phone=phone, email=cell(r, email_i), area=cell(r, area_i)))
    return finish_rows(rows)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1256", "replace")  # Arabic Windows exports


def parse_csv(data: bytes) -> list[Row]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return parse_table(list(csv.reader(io.StringIO(text), dialect)))


def parse_xlsx(data: bytes) -> list[Row]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    table = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i > MAX_ROWS + 1:
            break
        table.append(["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)) for v in r])
    wb.close()
    return parse_table(table)


# ---------------------------------------------------------------- pasted text
def parse_text(text: str) -> list[Row]:
    if "BEGIN:VCARD" in text.upper():
        return parse_vcf(text)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if lines and ("\t" in lines[0] or lines[0].count(",") >= 1) and any(k in lines[0].lower() for k in ("name", "phone", "mobile")):
        return parse_csv(text.encode())
    rows = []
    for ln in lines:
        ln = re.sub(r"^\[[^\]]*\]\s*", "", ln)  # WhatsApp-style "[10/10/26, 9:01 AM]"
        email_m = _EMAIL_RE.search(ln)
        email = email_m.group(0) if email_m else None
        if email:
            ln = ln.replace(email, " ")
        m = _PHONE_RE.search(ln)
        phone = clean_phone(m.group(1)) if m else None
        name = (ln[: m.start()] + " " + ln[m.end():]) if m else ln
        name = re.sub(r"[\s,;:\-–|]+", " ", name).strip(" -:")
        rows.append(Row(name=name, phone=phone, email=email))
    return finish_rows(rows)


def parse_file(filename: str, data: bytes) -> list[Row]:
    fn = (filename or "").lower()
    if fn.endswith(".vcf") or fn.endswith(".vcard") or data[:20].upper().lstrip().startswith(b"BEGIN:VCARD"):
        return parse_vcf(_decode(data))
    if fn.endswith(".xlsx") or data[:2] == b"PK":
        return parse_xlsx(data)
    if fn.endswith(".xls"):
        raise ValueError("Old .xls files aren't supported. In Excel, use File → Save As → .xlsx or CSV, then upload again.")
    return parse_csv(data)
